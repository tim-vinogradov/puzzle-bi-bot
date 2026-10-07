"""Выгрузка: таблица с вычислениями (CSV, Excel) и дашборд (интерактивный HTML)."""
from __future__ import annotations

import html
import io
import json
from datetime import datetime

import numpy as np
import pandas as pd

from . import fmt
from .calc import materialize
from .config import BASE_DIR
from .dashboard import DashboardState, render
from .semantic import DIM_LABELS, column_label, column_unit


def table_with_calcs(state: DashboardState) -> pd.DataFrame:
    frame = state.filtered(state.table)
    data = materialize(frame, state.calcs, dims=state.dim_columns(), measures=state.spec.measures,
                       prev=state.filtered_prev())
    data = data[[c for c in data.columns if c != "pct_sum"]].copy()
    if "period" in data.columns:
        data["period"] = pd.to_datetime(data["period"]).dt.strftime("%Y-%m")
    for c in data.columns:
        if pd.api.types.is_float_dtype(data[c]):
            data[c] = data[c].round(4 if column_unit(c, state.calcs) == "%" else 2)
    data.columns = [column_label(c, state.calcs) for c in data.columns]
    return data


def to_csv(state: DashboardState) -> bytes:
    data = table_with_calcs(state)
    return data.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig")


def to_xlsx(state: DashboardState) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = table_with_calcs(state)
    calc_labels = {c["label"] for c in state.calcs}
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        data.to_excel(writer, sheet_name="Данные", index=False)
        calcs = pd.DataFrame([{"Колонка": c["label"], "Формула": c.get("formula", ""), "Тип": c["kind"],
                               "Добавлена": c.get("created_at", "").replace("T", " ")} for c in state.calcs]
                             or [{"Колонка": "—", "Формула": "Вычислений нет", "Тип": "", "Добавлена": ""}])
        calcs.to_excel(writer, sheet_name="Вычисления", index=False)
        params = [
            ("Дашборд", state.name), ("Запрос", state.query), ("Набор данных", state.spec.label),
            ("Период", f"{state.period.label} ({state.period.range_label()})"),
            ("Фильтры", "; ".join(f"{DIM_LABELS.get(k, k)}: {', '.join(v)}" for k, v in state.filters.items() if v) or "нет"),
            ("Разрез", DIM_LABELS.get(state.switches.get("dim", ""), "")),
            ("Выгружено", datetime.now().strftime("%d.%m.%Y %H:%M")),
            ("SQL", state.sql),
        ]
        pd.DataFrame(params, columns=["Параметр", "Значение"]).to_excel(writer, sheet_name="Параметры", index=False)

        header_fill = PatternFill("solid", fgColor="EEF3FB")
        calc_fill = PatternFill("solid", fgColor="FFF4D6")
        for ws in writer.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.freeze_panes = "A2"
            for idx, col in enumerate(ws.columns, 1):
                values = [str(c.value) for c in list(col)[:200] if c.value is not None]
                width = min(60, max(10, max((len(v) for v in values), default=10) + 2))
                ws.column_dimensions[get_column_letter(idx)].width = width
        ws = writer.book["Данные"]
        for idx, name in enumerate(data.columns, 1):
            letter = get_column_letter(idx)
            if name in calc_labels:
                ws[f"{letter}1"].fill = calc_fill
            if pd.api.types.is_numeric_dtype(data[name]):
                number_format = "0.0" if name.endswith("%") else "# ##0.00" if pd.api.types.is_float_dtype(data[name]) else "# ##0"
                for cell in ws[letter][1:]:
                    cell.number_format = number_format
        ws.auto_filter.ref = ws.dimensions
    return buf.getvalue()


def to_html(state: DashboardState) -> bytes:
    payload = render(state)
    plotly_js = (BASE_DIR / "static" / "vendor" / "plotly.min.js").read_text(encoding="utf-8")
    locale_js = (BASE_DIR / "static" / "plotly-locale-ru.js").read_text(encoding="utf-8")
    esc = lambda s: html.escape(str(s))  # noqa: E731
    kpis = "".join(
        f'<div class="kpi {k["status"]}"><div class="kl">{esc(k["label"])}</div>'
        f'<div class="kv">{esc(k["arrow"] + " " if k["arrow"] else "")}{esc(k["display"])}</div>'
        f'<div class="ks">{esc(k["sub"])}</div></div>' for k in payload["kpis"])
    filters = "; ".join(f"{esc(DIM_LABELS.get(k, k))}: {esc(', '.join(v[:5]))}{'…' if len(v) > 5 else ''}"
                        for k, v in payload["active_filters"].items()) or "без фильтров"
    summary = "".join(f"<li>{s}</li>" for s in payload["summary"])
    charts_html = "".join(f'<div class="card"><div id="c{i}" class="chart"></div></div>' for i in range(len(payload["charts"])))
    figures = json.dumps([c["figure"] for c in payload["charts"]], ensure_ascii=False)
    calcs = "".join(f"<li><b>{esc(c['label'])}</b> = {esc(c['formula'])}</li>" for c in payload["calcs"])

    # сводная таблица по текущему разрезу
    table_html = ""
    try:
        from .calc import aggregate

        dim = state.switches.get("dim")
        frame = state.filtered(state.table)
        if dim in frame.columns:
            needed = [m for m in state.spec.measures if m in frame.columns and m != "pct_sum"] + [c["id"] for c in state.calcs]
            agg = aggregate(frame, [dim], needed, state.calcs, measures=state.spec.measures, prev=state.filtered_prev())
            agg = agg.sort_values(needed[0], ascending=False) if needed else agg
            head = "".join(f"<th>{esc(column_label(c, state.calcs))}</th>" for c in [dim] + needed)
            rows = []
            for _, row in agg.head(300).iterrows():
                cells = [f"<td>{esc(row[dim])}</td>"]
                for c in needed:
                    v = row[c]
                    unit = column_unit(c, state.calcs)
                    text = "—" if v is None or (isinstance(v, float) and np.isnan(v)) else (fmt.pct(v) if unit == "%" else fmt.num(v, 0))
                    cells.append(f'<td class="n">{esc(text)}</td>')
                rows.append("<tr>" + "".join(cells) + "</tr>")
            table_html = f'<details><summary>Таблица по разрезу «{esc(DIM_LABELS.get(dim, dim))}»</summary><div class="tw"><table><thead><tr>{head}</tr></thead><tbody>{"".join(rows)}</tbody></table></div></details>'
    except Exception:  # noqa: BLE001
        table_html = ""

    doc = f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(payload['name'] or 'Дашборд')}</title>
<style>
:root{{--bg:#f9f9f7;--card:#fcfcfb;--ink:#0b0b0b;--ink2:#52514e;--muted:#898781;--line:#e1e0d9;--good:#006300;--bad:#d03b3b}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--ink);font:14px/1.45 system-ui,-apple-system,"Segoe UI",Roboto,Arial,sans-serif}}
.wrap{{max-width:1320px;margin:0 auto;padding:24px 16px 48px}}h1{{font-size:22px;margin:0 0 4px}}.sub{{color:var(--ink2);margin-bottom:16px}}
.kpis{{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:10px;margin-bottom:14px}}
.kpi{{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:12px 14px}}.kl{{color:var(--ink2);font-size:12px}}
.kv{{font-size:22px;font-weight:650;margin-top:2px}}.ks{{color:var(--muted);font-size:12px}}.kpi.good .kv{{color:var(--good)}}.kpi.bad .kv{{color:var(--bad)}}
.grid{{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px}}@media(max-width:900px){{.grid{{grid-template-columns:1fr}}}}
.card{{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:6px}}.chart{{height:380px}}
.box{{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:12px 16px;margin:12px 0}}ul{{margin:6px 0;padding-left:20px}}
details{{margin-top:12px}}summary{{cursor:pointer;color:var(--ink2)}}.tw{{overflow:auto;max-height:520px;margin-top:8px}}
table{{border-collapse:collapse;width:100%;font-size:12.5px;font-variant-numeric:tabular-nums}}th,td{{border-bottom:1px solid var(--line);padding:5px 8px;text-align:left;white-space:nowrap}}
td.n{{text-align:right}}th{{position:sticky;top:0;background:#eef3fb}}.foot{{color:var(--muted);font-size:12px;margin-top:18px}}
</style></head><body><div class="wrap">
<h1>{esc(payload['name'] or 'Дашборд')}</h1>
<div class="sub">{esc(payload['dataset_label'])} · {esc(payload['period']['label'])} ({esc(payload['period_label'])}) · {filters}</div>
<div class="kpis">{kpis}</div>
<div class="box"><b>Главное</b><ul>{summary}</ul></div>
<div class="grid">{charts_html}</div>
{f'<div class="box"><b>Вычисления</b><ul>{calcs}</ul></div>' if calcs else ''}
{table_html}
<div class="foot">Запрос: «{esc(state.query)}». Срез на {datetime.now().strftime('%d.%m.%Y %H:%M')} с учётом фильтров. Графики интерактивные: зум (выделите область), подсказки, клик по легенде скрывает ряд, двойной клик — сброс.</div>
</div>
<script>{plotly_js}</script><script>{locale_js}</script>
<script>
const figs = {figures};
figs.forEach((f, i) => {{ f.layout.autosize = true; Plotly.newPlot('c' + i, f.data, f.layout, {{responsive: true, locale: 'ru', displaylogo: false}}); }});
</script></body></html>"""
    return doc.encode("utf-8")
